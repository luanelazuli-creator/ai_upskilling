"""Testes do pipeline de ingestão RAG (SPEC-004).

Usa o embedder "fake" (determinístico, sem download) para rodar rápido e offline.
Valida: chunking adaptativo, frontmatter por domínio, hash/re-indexação e
política de PII por domínio.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from src.rag.embedder import get_embedding_function
from src.rag.ingest import IngestionPipeline
from src.rag.preprocessing import build_frontmatter, profile_document
from src.rag.store import VectorStore


# --- fixtures ------------------------------------------------------------


def _make_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(path))


def _make_xlsx(path: Path, headers: list[str], rows: list[list]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(str(path))


@pytest.fixture
def data_tree(tmp_path: Path) -> Path:
    root = tmp_path / "data"
    _make_docx(
        root / "diario/2026/06/Dia 11 - Estudo.docx",
        "Dia 11 - Estudo",
        ["Estudei RAG.", "Liguei para (11) 98765-4321."],
    )
    _make_xlsx(
        root / "cursos/Cursos Feitos.xlsx",
        ["Curso", "Status"],
        [["DE Cert", "em_andamento"], ["ML Spec", "concluido"], ["Pydantic", "planejado"]],
    )
    _make_xlsx(
        root / "referencias/Pessoas Importantes.xlsx",
        ["Nome", "Cargo", "Email", "Telefone"],
        [["Ana Souza", "PM", "ana@x.com", "(11) 91234-5678"]],
    )
    return root


@pytest.fixture
def pipeline(tmp_path: Path) -> IngestionPipeline:
    store = VectorStore(
        db_path=str(tmp_path / "vs"),
        embedding_function=get_embedding_function("fake"),
    )
    return IngestionPipeline(
        store=store, processed_root=str(tmp_path / "processed")
    )


# --- chunking adaptativo (§4) --------------------------------------------


def test_diario_whole_file_one_chunk(data_tree, pipeline):
    report = pipeline.ingest_domain("diario", data_tree / "diario")
    assert report.total_chunks == 1  # 1 dia = 1 chunk


def test_xlsx_row_per_line(data_tree, pipeline):
    report = pipeline.ingest_domain("cursos", data_tree / "cursos")
    # 3 linhas na planilha => 3 chunks
    assert report.total_chunks == 3


# --- frontmatter (§3) ----------------------------------------------------


def test_frontmatter_diario_date_detection():
    path = Path("data/diario/2026/06/Dia 11 - Keep Studying.docx")
    profile = profile_document(path, "diario")
    fm = build_frontmatter(profile, path.name, "corpo")
    assert fm["type"] == "diario"
    assert fm["date"] == "2026-06-11"
    assert fm["content_hash"].startswith("sha256:")


def test_frontmatter_reference_category():
    path = Path("data/referencias/Pessoas Importantes.xlsx")
    profile = profile_document(path, "referencias")
    assert profile.category == "stakeholders"
    assert profile.pii_policy == "mask_contacts_only"


# --- hash / re-indexação (§6) --------------------------------------------


def test_unchanged_file_skipped(data_tree, pipeline):
    pipeline.ingest_domain("diario", data_tree / "diario")
    report2 = pipeline.ingest_domain("diario", data_tree / "diario")
    assert report2.files[0].status == "skipped"
    assert report2.total_chunks == 0


def test_modified_file_reindexed(data_tree, pipeline):
    pipeline.ingest_domain("diario", data_tree / "diario")
    # modifica o conteúdo do arquivo
    _make_docx(
        data_tree / "diario/2026/06/Dia 11 - Estudo.docx",
        "Dia 11 - Estudo",
        ["Conteúdo totalmente novo."],
    )
    report2 = pipeline.ingest_domain("diario", data_tree / "diario")
    assert report2.files[0].status == "reindexed"


def test_force_reindexes_unchanged(data_tree, pipeline):
    pipeline.ingest_domain("diario", data_tree / "diario")
    report2 = pipeline.ingest_domain("diario", data_tree / "diario", force=True)
    assert report2.files[0].status == "reindexed"


# --- política de PII (§7) ------------------------------------------------


def test_pii_aggressive_diario_redacts_phone(data_tree, pipeline):
    pipeline.ingest_domain("diario", data_tree / "diario")
    col = pipeline.store.get_or_create_collection("diario")
    doc = col.get(include=["documents"])["documents"][0]
    assert "(11) 98765-4321" not in doc
    assert "[REDACTED_PHONE]" in doc


def test_pii_stakeholders_keeps_name_masks_contacts(data_tree, pipeline):
    pipeline.ingest_domain("referencias", data_tree / "referencias")
    col = pipeline.store.get_or_create_collection("referencias")
    doc = col.get(include=["documents"])["documents"][0]
    assert "Ana Souza" in doc          # nome permanece
    assert "PM" in doc                  # cargo permanece
    assert "ana@x.com" not in doc       # email mascarado
    assert "[REDACTED_EMAIL]" in doc
    assert "[REDACTED_PHONE]" in doc


# --- busca + filtro estruturado (§8) -------------------------------------


def test_filter_by_category(data_tree, pipeline):
    pipeline.ingest_domain("referencias", data_tree / "referencias")
    results = pipeline.store.search(
        "referencias", "qualquer", top_k=5, where={"category": "stakeholders"}
    )
    assert len(results) >= 1
    assert all(r["metadata"]["category"] == "stakeholders" for r in results)
