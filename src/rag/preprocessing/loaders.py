"""Loaders por formato (SPEC-004 §2, etapa 2).

Cada loader converte um arquivo bruto em um RawDocument: corpo em Markdown
(para inspeção e para o .md processado) e, quando tabular, a lista de linhas
estruturadas usada pelo chunking row-per-line (SPEC-004 §4.1).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional


@dataclass
class RawDocument:
    source_path: Path
    body_markdown: str
    rows: Optional[List[Dict[str, Any]]] = None  # preenchido só para tabulares
    columns: List[str] = field(default_factory=list)


def load_document(path: Path) -> RawDocument:
    """Dispatcher por extensão."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".xlsx":
        return _load_xlsx(path)
    if suffix == ".csv":
        return _load_csv(path)
    if suffix == ".md":
        return _load_markdown(path)
    raise ValueError(f"Formato não suportado: {suffix}")


def _load_docx(path: Path) -> RawDocument:
    from docx import Document

    doc = Document(str(path))
    lines: List[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style else ""
        if style.startswith("Heading"):
            # "Heading 1" -> "# ", "Heading 2" -> "## "
            try:
                level = int(style.split()[-1])
            except (ValueError, IndexError):
                level = 2
            lines.append(f"{'#' * min(level, 6)} {text}")
        elif style == "Title":
            lines.append(f"# {text}")
        else:
            lines.append(text)
    return RawDocument(source_path=path, body_markdown="\n\n".join(lines))


def _rows_to_markdown_table(columns: List[str], rows: List[Dict[str, Any]]) -> str:
    if not columns:
        return ""
    header = "| " + " | ".join(columns) + " |"
    sep = "| " + " | ".join("---" for _ in columns) + " |"
    body_lines = []
    for row in rows:
        cells = [str(row.get(col, "")).replace("\n", " ").strip() for col in columns]
        body_lines.append("| " + " | ".join(cells) + " |")
    return "\n".join([header, sep, *body_lines])


def _normalize_rows(raw_rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Converte valores nulos/NaN em string vazia e força chaves string."""
    import math

    cleaned: List[Dict[str, Any]] = []
    for row in raw_rows:
        out: Dict[str, Any] = {}
        for key, value in row.items():
            if value is None:
                out[str(key)] = ""
            elif isinstance(value, float) and math.isnan(value):
                out[str(key)] = ""
            else:
                out[str(key)] = value
        cleaned.append(out)
    return cleaned


def _load_xlsx(path: Path) -> RawDocument:
    import pandas as pd

    df = pd.read_excel(path)
    columns = [str(c) for c in df.columns]
    rows = _normalize_rows(df.to_dict(orient="records"))
    body = _rows_to_markdown_table(columns, rows)
    return RawDocument(
        source_path=path, body_markdown=body, rows=rows, columns=columns
    )


def _load_csv(path: Path) -> RawDocument:
    import pandas as pd

    df = pd.read_csv(path)
    columns = [str(c) for c in df.columns]
    rows = _normalize_rows(df.to_dict(orient="records"))
    body = _rows_to_markdown_table(columns, rows)
    return RawDocument(
        source_path=path, body_markdown=body, rows=rows, columns=columns
    )


def _load_markdown(path: Path) -> RawDocument:
    text = path.read_text(encoding="utf-8")
    # Se já houver frontmatter, descarta — o pipeline gera o canônico.
    if text.startswith("---"):
        parts = text.split("---", 2)
        if len(parts) == 3:
            text = parts[2]
    return RawDocument(source_path=path, body_markdown=text.strip())
